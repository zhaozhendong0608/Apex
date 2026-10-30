# -*- coding: utf-8 -*-
"""把 pptx 按形状坐标画成 PNG，用于自查版式（不是交付渲染器）。"""

import sys
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import fontpath
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.dml import MSO_FILL

SCALE = 120.0  # px per inch
CJK = "/System/Library/Fonts/PingFang.ttc"
MONO = "/System/Library/Fonts/Menlo.ttc"
_CACHE = {}


def font(size_pt, bold=False, mono=False):
    px = max(1, int(round(size_pt * SCALE / 72.0)))
    key = (px, bold, mono)
    if key in _CACHE:
        return _CACHE[key]
    f = fontpath.load_mono(px) if mono else fontpath.load(px, bold)
    _CACHE[key] = f
    return f


def px(emu):
    return emu / 914400.0 * SCALE


def rgb(color_obj, default=(0, 0, 0)):
    try:
        c = color_obj.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return default


def draw_text_block(draw, shape, warn):
    tf = shape.text_frame
    ml, mt = px(tf.margin_left), px(tf.margin_top)
    mr, mb = px(tf.margin_right), px(tf.margin_bottom)
    x0, y0 = px(shape.left), px(shape.top)
    w, h = px(shape.width), px(shape.height)
    max_w = max(10, w - ml - mr)
    anchor = tf.vertical_anchor

    lines, line_hs, aligns = [], [], []
    for p in tf.paragraphs:
        runs = list(p.runs)
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        if not runs:
            lines.append([]); line_hs.append(0); aligns.append(p.alignment)
            continue
        max_size = max((r.font.size.pt if r.font.size else 14) for r in runs)
        lh = max_size * 1.22 * spacing * SCALE / 72
        cur, cur_w = [], 0.0
        for r in runs:
            size = r.font.size.pt if r.font.size else 14
            bold = bool(r.font.bold)
            mono = (r.font.name or "").lower().startswith("menlo")
            f = font(size, bold, mono)
            col = (0, 0, 0)
            try:
                if r.font.color is not None and r.font.color.type is not None:
                    col = rgb(r.font.color)
            except Exception:
                pass
            for ch in r.text:
                if ch == "\n":
                    lines.append(cur); line_hs.append(lh); aligns.append(p.alignment)
                    cur, cur_w = [], 0.0
                    continue
                cw = draw.textlength(ch, font=f)
                if cur_w + cw > max_w and cur:
                    lines.append(cur); line_hs.append(lh); aligns.append(p.alignment)
                    cur, cur_w = [], 0.0
                cur.append((ch, f, col))
                cur_w += cw
        lines.append(cur); line_hs.append(lh); aligns.append(p.alignment)

    total_h = sum(line_hs)
    if total_h > h + 2 and "".join(c for ln in lines for c, _, _ in ln).strip():
        warn.append((shape.name, round(h), round(total_h),
                     "".join(c for ln in lines for c, _, _ in ln)[:22]))

    if anchor == 3:
        y = y0 + (h - total_h) / 2
    elif anchor == 4:
        y = y0 + h - total_h - mb
    else:
        y = y0 + mt

    for ln, lh, al in zip(lines, line_hs, aligns):
        tw = sum(draw.textlength(c, font=f) for c, f, _ in ln)
        if al == 2:
            x = x0 + (w - tw) / 2
        elif al == 3:
            x = x0 + w - mr - tw
        else:
            x = x0 + ml
        for c, f, col in ln:
            draw.text((x, y), c, font=f, fill=col)
            x += draw.textlength(c, font=f)
        y += lh


def render(path, outdir):
    prs = Presentation(path)
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    W = int(prs.slide_width / 914400 * SCALE)
    H = int(prs.slide_height / 914400 * SCALE)
    warn = []
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), (247, 245, 240))
        draw = ImageDraw.Draw(img)
        try:
            bg = slide.background.fill.fore_color.rgb
            draw.rectangle([0, 0, W, H], fill=(bg[0], bg[1], bg[2]))
        except Exception:
            pass

        for shape in slide.shapes:
            tag = shape.element.tag
            if tag.endswith("}cxnSp"):
                col = rgb(shape.line.color, (200, 200, 200))
                wpt = shape.line.width.pt if shape.line.width else 1
                draw.line([px(shape.left), px(shape.top),
                           px(shape.left) + px(shape.width), px(shape.top) + px(shape.height)],
                          fill=col, width=max(1, int(wpt * SCALE / 72)))
                continue

            if getattr(shape, "has_table", False) and shape.has_table:
                tbl = shape.table
                yy = px(shape.top)
                for row in tbl.rows:
                    xx = px(shape.left)
                    rh = px(row.height)
                    for ci, cell in enumerate(row.cells):
                        cw = px(tbl.columns[ci].width)
                        try:
                            if cell.fill.type == MSO_FILL.SOLID:
                                fc = cell.fill.fore_color.rgb
                                draw.rectangle([xx, yy, xx + cw, yy + rh], fill=(fc[0], fc[1], fc[2]))
                        except Exception:
                            pass
                        # 表格单元格：手动定位
                        tf = cell.text_frame
                        for p in tf.paragraphs:
                            txt = "".join(r.text for r in p.runs)
                            if not txt:
                                continue
                            size = p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 11
                            bold = bool(p.runs[0].font.bold)
                            col = rgb(p.runs[0].font.color) if p.runs and p.runs[0].font.color.type is not None else (0, 0, 0)
                            f = font(size, bold)
                            tw = draw.textlength(txt, font=f)
                            lh = size * 1.22 * SCALE / 72
                            tx = xx + px(cell.margin_left)
                            if p.alignment == 2:
                                tx = xx + (cw - tw) / 2
                            ty = yy + (rh - lh) / 2
                            draw.text((tx, ty), txt, font=f, fill=col)
                        xx += cw
                    yy += rh
                continue

            # 普通形状
            fill = None
            try:
                if shape.fill.type == MSO_FILL.SOLID:
                    fill = rgb(shape.fill.fore_color)
            except Exception:
                fill = None

            outline = None
            lw = 1
            try:
                if shape.line.fill.type == MSO_FILL.SOLID:
                    outline = rgb(shape.line.color)
                    lw = max(1, int((shape.line.width.pt if shape.line.width else 1) * SCALE / 72))
            except Exception:
                outline = None

            x, y = px(shape.left), px(shape.top)
            w, h = px(shape.width), px(shape.height)
            kind = str(getattr(shape, "shape_type", ""))
            radius = 0
            if "ROUNDED" in kind:
                try:
                    radius = int(shape.adjustments[0] * min(w, h))
                except Exception:
                    radius = 0
            box = [x, y, x + w, y + h]
            if radius and fill:
                draw.rounded_rectangle(box, radius=radius, fill=fill,
                                       outline=outline, width=lw if outline else 0)
            elif fill or outline:
                draw.rectangle(box, fill=fill, outline=outline, width=lw if outline else 0)

            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_text_block(draw, shape, warn)

        img.save(outdir / f"slide-{idx:02d}.png")

    for w in warn:
        print("OVERFLOW?", w)
    print(f"rendered {len(prs.slides._sldIdLst)} slides -> {outdir}")


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
